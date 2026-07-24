"""文档解析模块。

本模块负责从不同格式的上传文件中提取纯文本内容，支持以下格式：
- 纯文本：.txt、.md、.markdown
- HTML：.html、.htm
- PDF：.pdf
- Word：.docx（旧版 .doc 不支持）

对于 HTML 和 DOCX，优先使用专业解析库（BeautifulSoup、python-docx），
当依赖不可用时回退到内置的 HTMLParser 兜底方案。
"""

from __future__ import annotations

from html.parser import HTMLParser
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile


# 支持的文件扩展名集合
SUPPORTED_EXTENSIONS = {".txt", ".md", ".markdown", ".html", ".htm", ".pdf", ".docx", ".doc"}


class KnowledgeParseError(ValueError):
    """文档解析失败异常。"""


def extract_text(filename: str, content: bytes) -> tuple[str, str]:
    """根据文件扩展名提取纯文本内容。

    根据文件后缀分派到对应的解析函数：
    - .txt/.md/.markdown：直接解码。
    - .html/.htm：解析 HTML 提取文本。
    - .pdf：使用 pypdf 提取文本。
    - .docx：使用 python-docx 或 ZIP XML 解析提取文本。
    - .doc：不支持，抛出异常。

    Args:
        filename: 文件名（含扩展名）。
        content: 文件二进制内容。

    Returns:
        ``(text, file_type)`` 元组，text 为提取的纯文本，file_type 为文件类型标识。

    Raises:
        KnowledgeParseError: 当文件格式不支持、缺少依赖库或解析失败时抛出。
    """
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise KnowledgeParseError(f"暂不支持 {suffix or 'unknown'} 文件格式。")
    if suffix == ".doc":
        raise KnowledgeParseError("暂不支持旧版 .doc 二进制格式，请转换为 .docx 后上传。")
    if suffix in {".txt", ".md", ".markdown"}:
        return _decode_text(content), suffix.lstrip(".")
    if suffix in {".html", ".htm"}:
        return _extract_html(content), "html"
    if suffix == ".pdf":
        return _extract_pdf(content), "pdf"
    if suffix == ".docx":
        return _extract_docx(content), "docx"
    raise KnowledgeParseError(f"暂不支持 {suffix} 文件格式。")


def _decode_text(content: bytes) -> str:
    """尝试用多种编码解码二进制内容为文本。

    依次尝试 UTF-8、UTF-8-SIG、GB18030、Latin-1，全部失败则忽略错误用 UTF-8 解码。
    """
    for encoding in ("utf-8", "utf-8-sig", "gb18030", "latin-1"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="ignore")


def _extract_html(content: bytes) -> str:
    """从 HTML 内容中提取纯文本。

    优先使用 BeautifulSoup（移除 script/style/noscript 标签后提取文本）；
    若依赖不可用则回退到内置的 _HTMLTextExtractor。
    """
    text = _decode_text(content)
    try:
        from bs4 import BeautifulSoup

        soup = BeautifulSoup(text, "html.parser")
        for item in soup(["script", "style", "noscript"]):
            item.decompose()
        return soup.get_text("\n")
    except Exception:
        parser = _HTMLTextExtractor()
        parser.feed(text)
        return parser.text


def _extract_pdf(content: bytes) -> str:
    """从 PDF 内容中提取纯文本。

    使用 pypdf 逐页提取文本，每页以 ``[Page N]`` 标记分隔。
    """
    try:
        from pypdf import PdfReader
    except Exception as exc:  # pragma: no cover - dependency availability differs by env.
        raise KnowledgeParseError("缺少 pypdf，无法解析 PDF。") from exc
    reader = PdfReader(BytesIO(content))
    pages: list[str] = []
    for index, page in enumerate(reader.pages):
        page_text = page.extract_text() or ""
        if page_text.strip():
            pages.append(f"[Page {index + 1}]\n{page_text}")
    return "\n\n".join(pages)


def _extract_docx(content: bytes) -> str:
    """从 DOCX 内容中提取纯文本。

    优先使用 python-docx 提取段落和表格文本；
    若依赖不可用则回退到 _extract_docx_with_zip 直接解析 XML。
    """
    try:
        from docx import Document

        document = Document(BytesIO(content))
        rows = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    rows.append(" | ".join(cells))
        return "\n".join(rows)
    except Exception:
        return _extract_docx_with_zip(content)


def _extract_docx_with_zip(content: bytes) -> str:
    """通过直接解析 DOCX 内部的 word/document.xml 提取文本（兜底方案）。

    当 python-docx 不可用时使用此方法，通过 ZipFile 读取 XML 并用
    _DocxTextExtractor 提取文本节点。
    """
    try:
        with ZipFile(BytesIO(content)) as archive:
            xml = archive.read("word/document.xml").decode("utf-8", errors="ignore")
    except Exception as exc:
        raise KnowledgeParseError("无法解析 docx 文档。") from exc
    parser = _DocxTextExtractor()
    parser.feed(xml)
    return parser.text


class _HTMLTextExtractor(HTMLParser):
    """内置 HTML 文本提取器（BeautifulSoup 不可用时的兜底方案）。

    通过 HTMLParser 收集所有非空文本节点，用换行符连接。
    """

    def __init__(self) -> None:
        super().__init__()
        self._parts: list[str] = []

    @property
    def text(self) -> str:
        return "\n".join(part.strip() for part in self._parts if part.strip())

    def handle_data(self, data: str) -> None:
        if data.strip():
            self._parts.append(data)


class _DocxTextExtractor(HTMLParser):
    """DOCX XML 文本提取器（python-docx 不可用时的兜底方案）。

    通过 HTMLParser 解析 word/document.xml，收集所有文本节点。
    """

    def __init__(self) -> None:
        super().__init__()
        self._parts: list[str] = []

    @property
    def text(self) -> str:
        return "\n".join(part.strip() for part in self._parts if part.strip())

    def handle_data(self, data: str) -> None:
        if data.strip():
            self._parts.append(data)
